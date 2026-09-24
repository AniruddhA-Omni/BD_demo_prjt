"""Structure-aware parsers that turn documents into ordered sections for chunking."""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Section:
    paragraphs: list[str] = field(default_factory=list)
    page: int | None = None
    section: str | None = None
    source_type: str = "text"


_HEADING_SIZE_RATIO = 1.15
_MAX_HEADING_CHARS = 120
_BOLD_FLAG = 16


def parse_pdf(path: Path) -> list[Section]:
    """Layout-aware PDF parsing: reading-order blocks, font-size/bold heading detection, OCR for scanned pages."""
    import pymupdf

    sections: list[Section] = []
    with pymupdf.open(path) as doc:
        pages = [page.get_text("dict", sort=True) for page in doc]
        body_size = _body_font_size(pages)
        current_heading: str | None = None

        for page_number, (page, layout) in enumerate(zip(doc, pages), start=1):
            section = Section(page=page_number, section=current_heading)
            has_text = False

            for block in layout.get("blocks", []):
                if block.get("type") != 0:
                    continue
                text, max_size, all_bold = _block_text(block)
                if not text:
                    continue
                has_text = True

                if _looks_like_heading(text, max_size, all_bold, body_size):
                    if section.paragraphs:
                        sections.append(section)
                    current_heading = text
                    section = Section(page=page_number, section=current_heading)
                else:
                    section.paragraphs.append(text)

            if not has_text:
                ocr_text = _ocr_pdf_page(page)
                if ocr_text:
                    section = Section(page=page_number, section=current_heading, source_type="ocr")
                    section.paragraphs.extend(_split_blocks(ocr_text))

            if section.paragraphs:
                sections.append(section)
    return sections


def parse_docx(path: Path) -> list[Section]:
    """DOCX parsing in document order; Heading/Title styles start new sections, tables become table sections."""
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    sections: list[Section] = []
    current = Section()

    for item in document.iter_inner_content():
        if isinstance(item, Table):
            rows = _table_rows([[cell.text for cell in row.cells] for row in item.rows])
            if rows:
                if current.paragraphs:
                    sections.append(current)
                sections.append(Section(paragraphs=["\n".join(rows)], section=current.section, source_type="table"))
                current = Section(section=current.section)
            continue

        text = item.text.strip()
        if not text:
            continue
        style_name = (item.style.name if item.style is not None else "") or ""
        if style_name.startswith("Heading") or style_name == "Title":
            if current.paragraphs:
                sections.append(current)
            current = Section(section=text)
        else:
            current.paragraphs.append(text)

    if current.paragraphs:
        sections.append(current)
    return sections


def parse_pptx(path: Path) -> list[Section]:
    """One section per slide (page = slide number, section = slide title), including tables and speaker notes."""
    from pptx import Presentation

    presentation = Presentation(str(path))
    sections: list[Section] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        paragraphs: list[str] = []

        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip()).strip()
                if text:
                    paragraphs.append(text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = _table_rows([[cell.text for cell in row.cells] for row in shape.table.rows])
                if rows:
                    paragraphs.append("\n".join(rows))

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                paragraphs.append(f"Speaker notes: {notes}")

        if not paragraphs and title:
            paragraphs.append(title)
        if paragraphs:
            sections.append(Section(paragraphs=paragraphs, page=slide_number, section=title or f"Slide {slide_number}"))
    return sections


def parse_markdown(text: str) -> list[Section]:
    sections: list[Section] = []
    current = Section()
    buffer: list[str] = []

    def flush() -> None:
        current.paragraphs.extend(_split_blocks("\n".join(buffer)))
        buffer.clear()

    for line in text.splitlines():
        heading = re.match(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$", line)
        if heading:
            flush()
            if current.paragraphs:
                sections.append(current)
            current = Section(section=heading.group(1) or None, source_type="markdown")
        else:
            buffer.append(line)
    flush()
    if current.paragraphs:
        current.source_type = "markdown"
        sections.append(current)
    return sections


def _split_blocks(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]


def _table_rows(rows: list[list[str]]) -> list[str]:
    rendered = []
    for row in rows:
        cells = [cell.strip() for cell in row]
        if any(cells):
            rendered.append(" | ".join(cells))
    return rendered


def _block_text(block: dict) -> tuple[str, float, bool]:
    lines: list[str] = []
    max_size = 0.0
    all_bold = True
    for line in block.get("lines", []):
        spans = [span for span in line.get("spans", []) if span.get("text", "").strip()]
        if not spans:
            continue
        lines.append("".join(span["text"] for span in line["spans"]).strip())
        for span in spans:
            max_size = max(max_size, float(span.get("size", 0.0)))
            all_bold = all_bold and bool(int(span.get("flags", 0)) & _BOLD_FLAG)
    text = " ".join(lines).strip()
    return text, max_size, all_bold and bool(lines)


def _body_font_size(pages: list[dict]) -> float:
    sizes: Counter[float] = Counter()
    for layout in pages:
        for block in layout.get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "").strip()
                    if text:
                        sizes[round(float(span.get("size", 0.0)), 1)] += len(text)
    return sizes.most_common(1)[0][0] if sizes else 0.0


def _looks_like_heading(text: str, size: float, bold: bool, body_size: float) -> bool:
    if len(text) > _MAX_HEADING_CHARS or text.endswith((".", ",", ";", ":")):
        return False
    if body_size and size >= body_size * _HEADING_SIZE_RATIO:
        return True
    return bold and len(text.split()) <= 12


def _ocr_pdf_page(page) -> str:
    try:
        import pytesseract
        from PIL import Image

        pixmap = page.get_pixmap(dpi=200)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        return pytesseract.image_to_string(image).strip()
    except Exception:
        return ""
