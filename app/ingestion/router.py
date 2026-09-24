from __future__ import annotations

import logging
import time
from pathlib import Path
from typing import Any
import uuid

from app.ingestion.chunking import chunk_paragraphs, split_paragraphs
from app.ingestion.code_parser import parse_python
from app.ingestion.models import Evidence
from app.ingestion.parsers import Section, parse_docx, parse_markdown, parse_pdf, parse_pptx
from app.ingestion.tables import TableSection, read_table_sections


logger = logging.getLogger(__name__)

VISION_DESCRIPTION_PROMPT = (
    "Describe this image for a document search index. State what kind of image it is (chart, diagram, screenshot, "
    "photo, scanned page, form). Transcribe titles, labels, axis names and legend entries; for charts give the key "
    "values and trends; for diagrams list the components and how they connect. Only describe what is visible."
)


def _cell_range(section: TableSection) -> str | None:
    """Spreadsheet range covered by a table, header row included, e.g. ``A3:B6``."""
    from openpyxl.utils import get_column_letter

    rows, columns = section.frame.shape
    if not columns:
        return None
    last_row = section.row_numbers[-1] if section.row_numbers else section.header_row + rows
    first = get_column_letter(section.first_col)
    last = get_column_letter(section.first_col + columns - 1)
    return f"{first}{section.header_row}:{last}{last_row}"


class IngestionRouter:
    """Dispatches files to format-specific parsers and emits chunked, citation-ready evidence."""

    _TEXT_TYPES = {"txt", "md"}
    _MARKDOWN_TYPES = {"md"}
    _DOCUMENT_PARSERS = {"pdf": parse_pdf, "docx": parse_docx, "pptx": parse_pptx}
    _TABLE_TYPES = {"csv", "xlsx"}
    _IMAGE_TYPES = {"png", "jpg", "jpeg"}
    _CODE_TYPES = {"py"}

    def ingest_file(self, file_path: str | Path) -> list[Evidence]:
        path = Path(file_path)
        file_type = path.suffix.lower().lstrip(".") or "unknown"
        doc_id = str(uuid.uuid4())

        started = time.perf_counter()
        try:
            evidence = self._dispatch(path, doc_id, file_type)
        except Exception as exc:
            logger.exception("Failed to parse %s (%s)", path.name, file_type)
            return [
                Evidence(
                    document_id=doc_id,
                    file_name=path.name,
                    file_type=file_type,
                    source_type="error",
                    content=f"Could not parse {path.name}: {exc}",
                    metadata=self._stamp_metadata(path, status="failed", source_type="error"),
                )
            ]

        self._log_result(path.name, file_type, evidence, (time.perf_counter() - started) * 1000)
        return evidence

    def _log_result(self, name: str, file_type: str, evidence: list[Evidence], elapsed_ms: float) -> None:
        statuses = {str(item.metadata.get("status")) for item in evidence}
        if "parsed" not in statuses:
            logger.warning("Document not loaded: %s (%s) status=%s", name, file_type, ",".join(sorted(statuses)) or "empty")
            return
        parts = [f"{len(evidence)} chunk(s)"]
        pages = {item.page for item in evidence if item.page is not None}
        if pages:
            parts.append(f"{len(pages)} {'slide' if file_type == 'pptx' else 'page'}(s)")
        sheets = {item.sheet for item in evidence if item.sheet}
        if sheets:
            parts.append(f"{len(sheets)} sheet(s)")
        sections = {item.section for item in evidence if item.section}
        if sections and not sheets:
            parts.append(f"{len(sections)} section(s)")
        placeholders = sum(1 for item in evidence if item.metadata.get("placeholder"))
        if placeholders:
            parts.append(f"{placeholders} placeholder(s): no text extracted")
        logger.info("Document loaded successfully: %s (%s): %s in %.0f ms", name, file_type, ", ".join(parts), elapsed_ms)

    def _dispatch(self, path: Path, doc_id: str, file_type: str) -> list[Evidence]:
        if file_type in self._CODE_TYPES:
            return self._parse_code_file(path, doc_id, file_type)

        if file_type in self._TEXT_TYPES:
            return self._parse_text_file(path, doc_id, file_type)

        if file_type in self._TABLE_TYPES:
            return self._parse_table_file(path, doc_id, file_type)

        if file_type in self._IMAGE_TYPES:
            return self._parse_image_file(path, doc_id, file_type)

        if file_type in self._DOCUMENT_PARSERS:
            sections = self._DOCUMENT_PARSERS[file_type](path)
            evidence = self._chunk_sections(path, doc_id, file_type, sections)
            if evidence:
                return evidence
            return [
                Evidence(
                    document_id=doc_id,
                    file_name=path.name,
                    file_type=file_type,
                    source_type="empty",
                    content=f"No extractable text was found in {path.name}.",
                    metadata=self._stamp_metadata(path, status="empty", source_type="empty"),
                )
            ]

        return [
            Evidence(
                document_id=doc_id,
                file_name=path.name,
                file_type=file_type,
                source_type="unknown",
                content=f"Unsupported file type '{file_type}' queued for future handling.",
                metadata=self._stamp_metadata(path, status="unsupported", source_type="unknown"),
            )
        ]

    def _chunk_sections(self, path: Path, doc_id: str, file_type: str, sections: list[Section]) -> list[Evidence]:
        evidence: list[Evidence] = []
        for section in sections:
            # Tables are kept whole so rows are never split away from their siblings.
            if section.source_type == "table":
                chunks = ["\n\n".join(section.paragraphs)]
            else:
                chunks = chunk_paragraphs(section.paragraphs)
            for chunk in chunks:
                index = len(evidence)
                metadata = self._stamp_metadata(path, status="parsed", source_type=section.source_type)
                metadata["chunk_index"] = index
                evidence.append(
                    Evidence(
                        document_id=doc_id,
                        file_name=path.name,
                        file_type=file_type,
                        source_type=section.source_type,
                        page=section.page,
                        section=section.section,
                        chunk_id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_id}:{index}")),
                        content=chunk,
                        metadata=metadata,
                    )
                )
        for item in evidence:
            item.metadata["chunk_count"] = len(evidence)
        return evidence

    def _parse_text_file(self, path: Path, doc_id: str, file_type: str) -> list[Evidence]:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        source_type = "markdown" if file_type in self._MARKDOWN_TYPES else "text"
        if source_type == "markdown":
            sections = parse_markdown(text)
        else:
            sections = [Section(paragraphs=split_paragraphs(text), source_type="text")]

        evidence = self._chunk_sections(path, doc_id, file_type, sections)
        if evidence:
            return evidence
        return [
            Evidence(
                document_id=doc_id,
                file_name=path.name,
                file_type=file_type,
                source_type=source_type,
                content="Empty file.",
                metadata=self._stamp_metadata(path, status="parsed", source_type=source_type),
            )
        ]

    def _parse_code_file(self, path: Path, doc_id: str, file_type: str) -> list[Evidence]:
        source = path.read_text(encoding="utf-8", errors="ignore")
        evidence = []
        for index, chunk in enumerate(parse_python(source, path.name)):
            evidence.append(
                self._evidence(path, doc_id, file_type, index, "code", chunk.content, section=chunk.section, extra=chunk.metadata)
            )
        return evidence

    def _parse_table_file(self, path: Path, doc_id: str, file_type: str) -> list[Evidence]:
        """One ``table`` chunk per CSV / XLSX sheet: normalised CSV content plus where the table sits in the sheet.

        The content is clean CSV (header first) so the data agent can rebuild the DataFrame; ``header_row``,
        ``first_col`` and ``row_numbers`` let it cite exact cells even when the header isn't in row 1.
        """
        evidence = []
        for section in read_table_sections(path, file_type):
            extra: dict[str, Any] = {"header_row": section.header_row, "first_col": section.first_col}
            if section.row_numbers:
                extra["row_numbers"] = section.row_numbers
            if section.title:
                extra["title"] = section.title
            evidence.append(
                self._evidence(
                    path,
                    doc_id,
                    file_type,
                    len(evidence),
                    "table",
                    section.frame.to_csv(index=False).strip(),
                    section=section.title,
                    sheet=section.sheet,
                    cell_range=_cell_range(section),
                    extra=extra,
                )
            )
        return evidence or [self._evidence(path, doc_id, file_type, 0, "table", "Empty table.")]

    def _parse_image_file(self, path: Path, doc_id: str, file_type: str) -> list[Evidence]:
        ocr_status = "ok"
        try:
            from PIL import Image
            import pytesseract

            with Image.open(path) as image:
                extracted_text = pytesseract.image_to_string(image).strip()
            content = extracted_text or f"Image file '{path.name}' was processed, but no text was detected by OCR."
            if not extracted_text:
                ocr_status = "no_text"
        except Exception:
            ocr_status = "unavailable"
            content = (
                f"Image file '{path.name}' was ingested, but OCR is unavailable in this environment. "
                "Install a working Tesseract binary to extract text from images."
            )
        extra = {"ocr_status": ocr_status, "placeholder": ocr_status != "ok"}
        evidence = [self._evidence(path, doc_id, file_type, 0, "image", content, section="OCR text", extra=extra)]

        # Visual understanding is a separate step from OCR: layout, charts and diagrams carry meaning OCR can't read.
        # It is slow, so by default the vision agent describes an image on the first question about it (and caches
        # the description); VISION_AT_INGESTION=true does it here instead.
        from app import services
        from app.config import load_settings

        if load_settings().vision_at_ingestion:
            description = services.get_llm().describe_images(VISION_DESCRIPTION_PROMPT, [str(path)], run_name="vision_ingest")
            if description:
                evidence.append(self._evidence(path, doc_id, file_type, 1, "vision", description, section="visual description"))
        return evidence

    def _evidence(
        self,
        path: Path,
        doc_id: str,
        file_type: str,
        index: int,
        source_type: str,
        content: str,
        *,
        page: int | None = None,
        section: str | None = None,
        sheet: str | None = None,
        cell_range: str | None = None,
        extra: dict[str, Any] | None = None,
    ) -> Evidence:
        metadata = self._stamp_metadata(path, status="parsed", source_type=source_type)
        metadata["chunk_index"] = index
        metadata.update(extra or {})
        return Evidence(
            document_id=doc_id,
            file_name=path.name,
            file_type=file_type,
            source_type=source_type,
            page=page,
            section=section,
            sheet=sheet,
            cell_range=cell_range,
            chunk_id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_id}:{index}")),
            content=content,
            metadata=metadata,
        )

    def _stamp_metadata(self, path: Path, *, status: str, source_type: str) -> dict[str, Any]:
        return {
            "source_name": path.name,
            "source_path": str(path),
            "source_type": source_type,
            "path": str(path),
            "status": status,
        }
