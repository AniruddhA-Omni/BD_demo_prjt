"""Spreadsheet/CSV parsing with header-row detection.

Exported spreadsheets often have a title, notes or blank rows above the real header, or start in a column other
than A. Reading them with ``header=0`` produces "Unnamed" columns and cell ranges that point at the wrong cells.
This module reads the raw cell grid (openpyxl for XLSX, the csv module for CSV), finds the header row, and records
where the table sits so citations like ``sheet Budget — range A3:B6`` are exact.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

_HEADER_SCAN_ROWS = 15


@dataclass
class TableSection:
    frame: pd.DataFrame
    sheet: str | None
    header_row: int  # 1-based spreadsheet row holding the column names
    first_col: int  # 1-based column index of the table's first column
    title: str | None = None  # text found above the header (report titles, notes)
    row_numbers: list[int] | None = field(default=None)  # 1-based rows of each data row, only when not contiguous


def read_table_sections(path: Path, file_type: str) -> list[TableSection]:
    grids = _read_csv_grid(path) if file_type == "csv" else _read_xlsx_grids(path)
    sections = []
    for sheet, grid in grids:
        section = normalise_grid(grid, sheet)
        if section is not None:
            sections.append(section)
    return sections


def normalise_grid(grid: list[list[Any]], sheet: str | None = None) -> TableSection | None:
    """Turn a raw cell grid into a typed DataFrame plus its position in the sheet."""
    rows = [[_clean(value) for value in row] for row in grid]
    non_empty = [index for index, row in enumerate(rows) if any(value is not None for value in row)]
    if not non_empty:
        return None

    header_index = _find_header_row(rows, non_empty)
    header = rows[header_index]
    columns = [index for index, value in enumerate(header) if value is not None]
    first, last = columns[0], columns[-1]

    names = _unique_names([header[index] for index in range(first, last + 1)])
    data: list[list[Any]] = []
    row_numbers: list[int] = []
    for index in range(header_index + 1, len(rows)):
        values = (rows[index] + [None] * (last + 1))[first : last + 1]
        if any(value is not None for value in values):
            data.append(values)
            row_numbers.append(index + 1)
    if not data:
        return None

    frame = pd.DataFrame(data, columns=names)
    for name in frame.columns:
        frame[name] = _maybe_numeric(frame[name])

    contiguous = row_numbers == list(range(header_index + 2, header_index + 2 + len(row_numbers)))
    title_rows = [" ".join(str(v) for v in rows[i] if v is not None) for i in non_empty if i < header_index]
    return TableSection(
        frame=frame,
        sheet=sheet,
        header_row=header_index + 1,
        first_col=first + 1,
        title=" / ".join(title_rows) or None,
        row_numbers=None if contiguous else row_numbers,
    )


def _find_header_row(rows: list[list[Any]], non_empty: list[int]) -> int:
    """First row (within the scan window) that looks like column names and is followed by data.

    A header has at least two cells, all of them non-numeric text, and spans most of the columns the next rows use.
    Falls back to the first non-empty row.
    """
    for index in non_empty[:_HEADER_SCAN_ROWS]:
        cells = [value for value in rows[index] if value is not None]
        if len(cells) < 2 or not all(isinstance(value, str) and not _is_number(value) for value in cells):
            continue
        following = [i for i in non_empty if i > index][:5]
        if not following:
            continue
        widest = max(sum(1 for value in rows[i] if value is not None) for i in following)
        if len(cells) >= max(2, int(0.6 * widest)):
            return index
    return non_empty[0]


def _read_xlsx_grids(path: Path) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            (sheet.title, [list(row) for row in sheet.iter_rows(min_row=1, min_col=1, values_only=True)])
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _read_csv_grid(path: Path) -> list[tuple[None, list[list[Any]]]]:
    with path.open(newline="", encoding="utf-8", errors="ignore") as handle:
        return [(None, [row for row in csv.reader(handle)])]


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _is_number(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except ValueError:
        return False


def _maybe_numeric(series: pd.Series) -> pd.Series:
    """Convert a column to numbers only when every non-empty value is numeric (CSV cells arrive as strings)."""
    as_text = series.map(lambda v: v.replace(",", "") if isinstance(v, str) else v)
    converted = pd.to_numeric(as_text, errors="coerce")
    return converted if converted.notna().sum() == series.notna().sum() else series


def _unique_names(values: list[Any]) -> list[str]:
    names: list[str] = []
    for position, value in enumerate(values, start=1):
        name = str(value) if value is not None else f"column_{position}"
        candidate, suffix = name, 2
        while candidate in names:
            candidate, suffix = f"{name}_{suffix}", suffix + 1
        names.append(candidate)
    return names
