from pathlib import Path

import docx
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.agents.synthesis import format_citation
from app.ingestion.chunking import chunk_paragraphs
from app.ingestion.router import IngestionRouter
from app.retrieval.hybrid import HybridRetriever


def _make_pdf(path: Path) -> None:
    # PyMuPDF's native DLLs need the MSVC runtime on Windows; skip rather than fail where it is missing.
    pymupdf = pytest.importorskip("pymupdf", exc_type=ImportError)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Revenue Outlook", fontsize=20)
    page.insert_text((72, 110), "Revenue increased by 12% in Q4 driven by enterprise sales.", fontsize=11)
    page = doc.new_page()
    page.insert_text((72, 72), "Risks", fontsize=20)
    page.insert_text((72, 110), "Supply chain delays may affect hardware margins next year.", fontsize=11)
    doc.save(path)
    doc.close()


def test_pdf_parsing_tracks_pages_and_headings(tmp_path: Path):
    path = tmp_path / "annual_report.pdf"
    _make_pdf(path)

    evidence = IngestionRouter().ingest_file(path)

    assert [(e.page, e.section) for e in evidence] == [(1, "Revenue Outlook"), (2, "Risks")]
    assert "12%" in evidence[0].content
    assert all(e.metadata["status"] == "parsed" for e in evidence)
    assert len({e.chunk_id for e in evidence}) == 2
    assert len({e.document_id for e in evidence}) == 1
    assert format_citation(evidence[1]) == "annual_report.pdf — page 2 — Risks"


def test_docx_parsing_uses_heading_styles_and_tables(tmp_path: Path):
    path = tmp_path / "company_notes.docx"
    document = docx.Document()
    document.add_heading("Hiring Plan", level=1)
    document.add_paragraph("We plan to hire 40 engineers in 2027.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Team", "Headcount"
    table.cell(1, 0).text, table.cell(1, 1).text = "Platform", "12"
    document.save(str(path))

    evidence = IngestionRouter().ingest_file(path)

    assert [e.source_type for e in evidence] == ["text", "table"]
    assert all(e.section == "Hiring Plan" for e in evidence)
    assert "Platform | 12" in evidence[1].content


def test_pptx_parsing_emits_one_section_per_slide(tmp_path: Path):
    path = tmp_path / "presentation.pptx"
    presentation = Presentation()
    for title, body in [("Strategy", "Expand into EMEA markets."), ("Roadmap", "Launch v2 in March.")]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    slide.notes_slide.notes_text_frame.text = "Mention the beta customers."
    slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = "Owner: Platform team"
    presentation.save(str(path))

    evidence = IngestionRouter().ingest_file(path)

    assert [(e.page, e.section) for e in evidence] == [(1, "Strategy"), (2, "Roadmap")]
    assert "Speaker notes: Mention the beta customers." in evidence[1].content
    assert "Owner: Platform team" in evidence[1].content
    assert format_citation(evidence[0]) == "presentation.pptx — slide 1 — Strategy"


def test_markdown_headings_become_sections(tmp_path: Path):
    path = tmp_path / "notes.md"
    path.write_text("# Intro\nHello world.\n\n## Pricing\nPlans start at $10.\n", encoding="utf-8")

    evidence = IngestionRouter().ingest_file(path)

    assert [e.section for e in evidence] == ["Intro", "Pricing"]
    assert all(e.source_type == "markdown" for e in evidence)


def test_long_text_is_chunked_without_losing_content(tmp_path: Path):
    paragraphs = [f"Paragraph {i} discusses topic number {i} in some detail." for i in range(200)]
    path = tmp_path / "long.txt"
    path.write_text("\n\n".join(paragraphs), encoding="utf-8")

    evidence = IngestionRouter().ingest_file(path)

    assert len(evidence) > 1
    assert all(len(e.content) <= 1200 for e in evidence)
    joined = "\n\n".join(e.content for e in evidence)
    assert all(p in joined for p in paragraphs)
    assert [e.metadata["chunk_index"] for e in evidence] == list(range(len(evidence)))


def test_chunker_splits_single_oversized_paragraph():
    chunks = chunk_paragraphs(["word " * 1000], max_chars=300)

    assert len(chunks) > 1
    assert all(len(chunk) <= 300 for chunk in chunks)


def test_retrieval_ranks_chunks_of_the_same_document_independently(tmp_path: Path):
    path = tmp_path / "annual_report.pdf"
    _make_pdf(path)
    evidence = IngestionRouter().ingest_file(path)

    retriever = HybridRetriever()
    retriever.qdrant.enabled = False
    ranked = retriever.search("supply chain delays", evidence, top_k=2)

    assert len(ranked) == 2
    assert ranked[0].section == "Risks"
